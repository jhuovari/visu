#!/usr/bin/env python3
"""Kokoa esitys visu-sivuston kuvioista.

Lukee maarittelytiedoston (YAML), renderoi kuviot R:lla esityspohjan
paikkamerkin mittoihin ja kokoaa diat PowerPoint-tiedostoksi.

    python3 esitys/kokoa.py esitys/suhdanne.yml

Maarittelytiedoston muoto on kuvattu tiedostossa esitys/README.md.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import subprocess
import sys
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.util import Inches

JUURI = Path(__file__).resolve().parent.parent
OLETUSPOHJA = JUURI / "esitys" / "pohja.pptx"

# Pohjan asettelut nimella. Numero on asettelun jarjestysnumero masterissa.
ASETTELUT = {
    "kansi": 0,
    "kuvio": 6,          # "Kaksi palstaa tai sisalto ja graafi"
    "valiotsikko": 9,
    "vain_otsikko": 14,
    "lopetus": 17,
}


def paiva(arvo):
    """YAML lukee pp-muotoisen paivan date-olioksi; R haluaa merkkijonon."""
    if arvo is None:
        return None
    if isinstance(arvo, (dt.date, dt.datetime)):
        return arvo.strftime("%Y-%m-%d")
    return str(arvo)


def laske_mitat(layout, idx):
    """Paikkamerkin sijainti ja koko tuumina."""
    for ph in layout.placeholders:
        if ph.placeholder_format.idx == idx:
            return (ph.left / 914400, ph.top / 914400,
                    ph.width / 914400, ph.height / 914400)
    raise KeyError(f"Asettelussa {layout.name!r} ei ole paikkamerkkia {idx}")


def renderoi(kuviot, kansio, site_dir):
    """Renderoi kuviot R:lla yhdella ajolla.

    Yksi ajo per esitys eika per kuvio, koska sivun lataaminen hakee kaikkien
    sen kuvioiden datan rajapinnasta: saman sivun seuraavat kuviot ovat
    saman prosessin sisalla ilmaisia.
    """
    if not kuviot:
        return {}
    kansio.mkdir(parents=True, exist_ok=True)
    pyynnot = []
    for k in kuviot:
        tiedosto = kansio / f"{k['tiedosto']}.png"
        pyynnot.append({**k, "ulos": str(tiedosto)})

    skripti = r"""
suppressMessages(library(visu))
pyynnot <- jsonlite::fromJSON(commandArgs(TRUE)[1], simplifyDataFrame = FALSE)
site_dir <- commandArgs(TRUE)[2]
for (p in pyynnot) {
  visu_chart_png(
    p$id, p$ulos, lang = p$kieli,
    start = if (is.null(p$alkaen)) NULL else p$alkaen,
    end = if (is.null(p$asti)) NULL else p$asti,
    titles = isTRUE(p$otsikko_kuviossa),
    width = p$leveys, height = p$korkeus, site_dir = site_dir
  )
  cat("renderoitu:", p$ulos, "\n")
}
"""
    with (kansio / "_pyynnot.json").open("w", encoding="utf-8") as f:
        json.dump(pyynnot, f, ensure_ascii=False)
    aja = subprocess.run(
        ["Rscript", "-e", skripti, str(kansio / "_pyynnot.json"), str(site_dir)],
        cwd=JUURI, capture_output=True, text=True,
        env={"LANG": "C.utf8", "LC_ALL": "C.utf8", "PATH": "/usr/bin:/bin:/usr/local/bin",
             "HOME": str(Path.home())},
    )
    if aja.returncode != 0:
        sys.stderr.write(aja.stdout + aja.stderr)
        raise SystemExit("Kuvioiden renderointi epaonnistui.")
    return {k["tiedosto"]: Path(k["ulos"]) for k in pyynnot}


RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def tyhjenna(prs):
    """Poista pohjan mallidiat.

    Pelkka poisto dialuettelosta ei riita: dian osa jaa pakettiin ja
    tallennus kirjoittaisi sen uudelleen samalla nimella kuin uuden dian.
    Siksi myos suhde presentation-osaan katkaistaan, jolloin osa jaa
    orvoksi eika paady tiedostoon.
    """
    luettelo = prs.slides._sldIdLst
    for dia in list(luettelo):
        prs.part.drop_rel(dia.get(RID))
        luettelo.remove(dia)


def poista(muoto):
    """Irrota muoto dialta. Tyhja paikkamerkki nayttaisi kehotetekstin."""
    muoto._element.getparent().remove(muoto._element)


def paikkamerkki(dia, idx):
    for ph in dia.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def teksti(dia, idx, arvo):
    ph = paikkamerkki(dia, idx)
    if ph is None:
        return
    if arvo is None:
        poista(ph)
        return
    kehys = ph.text_frame
    rivit = arvo if isinstance(arvo, list) else [arvo]
    kehys.text = str(rivit[0])
    for rivi in rivit[1:]:
        kehys.add_paragraph().text = str(rivi)


# Koko leveyden kuviopaikka: [x, y, leveys, korkeus] tuumina, luetaan
# pohjasta kokoa():ssa.
LEVEA: list[float] = [0, 0, 0, 0]


def lisaa_kuvio(dia, idx, kuva, levea=False):
    """Kuva paikkamerkin tilalle sen omilla mitoilla.

    Sisaltopaikkamerkkiin ei voi python-pptx:lla pudottaa kuvaa suoraan, joten
    mitat luetaan paikkamerkista, paikkamerkki poistetaan ja kuva lisataan
    samaan kohtaan. Kuva on renderoitu taman kokoiseksi, joten se tayttaa
    paikan ilman venytysta.
    """
    ph = paikkamerkki(dia, idx)
    if ph is None:
        raise KeyError(f"Dialla ei ole paikkamerkkia {idx}")
    if levea:
        vasen, ylos, leveys, korkeus = (Inches(v) for v in LEVEA)
    else:
        vasen, ylos, leveys, korkeus = ph.left, ph.top, ph.width, ph.height
    poista(ph)
    dia.shapes.add_picture(str(kuva), vasen, ylos, width=leveys, height=korkeus)


def kokoa(maarittely: Path, ulos: Path | None = None):
    with maarittely.open(encoding="utf-8") as f:
        m = yaml.safe_load(f)

    # Liitettaessa olemassa olevaan esitykseen sen omat diat jaavat paikalleen
    # ja uudet tulevat peraan; omasta pohjasta mallidiat poistetaan.
    liitetaan = bool(m.get("liita"))
    lahde = Path(m.get("liita") or m.get("pohja") or OLETUSPOHJA)
    if not lahde.is_absolute():
        lahde = JUURI / lahde
    prs = Presentation(str(lahde))
    master = prs.slide_masters[0]
    if not liitetaan:
        tyhjenna(prs)
    asettelu = master.slide_layouts[ASETTELUT["kuvio"]]
    vasen_x, _, _, _ = laske_mitat(asettelu, 1)
    oikea_x, oikea_y, kuva_w, kuva_h = laske_mitat(asettelu, 2)
    # Ilman tekstipalstaa kuvio saa molempien palstojen leveyden. Mitat
    # tulevat yha pohjasta: vasen reuna vasemmasta palstasta, oikea reuna
    # oikeasta, korkeus kuviopalstasta.
    leve_w = oikea_x + kuva_w - vasen_x
    LEVEA[:] = [vasen_x, oikea_y, leve_w, kuva_h]

    kieli = m.get("kieli", "fi")
    site_dir = m.get("site_dir", "site")
    tyokansio = maarittely.parent / "kuvat" / maarittely.stem

    pyynnot = []
    for i, dia in enumerate(m.get("diat", [])):
        if "kuvio" not in dia:
            continue
        levea = not dia.get("teksti")
        pyynnot.append({
            "id": dia["kuvio"],
            "tiedosto": f"{i:02d}-{dia['kuvio']}",
            "kieli": dia.get("kieli", kieli),
            "alkaen": paiva(dia.get("alkaen")),
            "asti": paiva(dia.get("asti")),
            "otsikko_kuviossa": dia.get("otsikko_kuviossa", True),
            "leveys": round(LEVEA[2] if levea else kuva_w, 3),
            "korkeus": round(kuva_h, 3),
        })
    kuvat = renderoi(pyynnot, tyokansio, site_dir)

    for i, dia in enumerate(m.get("diat", [])):
        if "kuvio" in dia:
            s = prs.slides.add_slide(master.slide_layouts[ASETTELUT["kuvio"]])
            teksti(s, 0, dia.get("otsikko", ""))
            teksti(s, 1, dia.get("teksti"))
            lisaa_kuvio(s, 2, kuvat[f"{i:02d}-{dia['kuvio']}"],
                        levea=not dia.get("teksti"))
            if dia.get("muistiinpanot"):
                s.notes_slide.notes_text_frame.text = dia["muistiinpanot"]
        elif "kansi" in dia:
            s = prs.slides.add_slide(master.slide_layouts[ASETTELUT["kansi"]])
            teksti(s, 0, dia["kansi"])
            teksti(s, 1, dia.get("alaotsikko"))
        elif "valiotsikko" in dia:
            s = prs.slides.add_slide(master.slide_layouts[ASETTELUT["valiotsikko"]])
            teksti(s, 0, dia["valiotsikko"])
        elif "lopetus" in dia:
            s = prs.slides.add_slide(master.slide_layouts[ASETTELUT["lopetus"]])
            teksti(s, 0, dia["lopetus"])
            teksti(s, 1, dia.get("alaotsikko"))
        else:
            raise ValueError(f"Dialta {i} puuttuu tyyppi (kuvio, kansi, valiotsikko, lopetus).")

    kohde = ulos or (maarittely.parent / f"{maarittely.stem}.pptx")
    prs.save(str(kohde))
    print(f"valmis: {kohde} ({len(prs.slides)} diaa)")
    return kohde


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("maarittely", type=Path)
    ap.add_argument("-o", "--ulos", type=Path, default=None)
    kokoa(**vars(ap.parse_args()))
